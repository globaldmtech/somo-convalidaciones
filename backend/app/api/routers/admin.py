import io
import os
import sqlite3
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment

from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response
from fastapi.responses import FileResponse
from fastapi.routing import APIRoute

from ...auth.session_user import require_admin_user
from ...db import database
from ...model.admin import (
    ActualizarCicloRequest,
    ActualizarUsuarioRolRequest,
    CambiarEstadoFormularioRequest,
    CambiarEstadoSolicitudRequest,
    EliminarConvalidacionesMasivasRequest,
    EliminarConvalidacionesCicloMasivasRequest,
    CrearCicloRequest,
    CrearConvalidacionRequest,
    CrearConvalidacionesMasivasRequest,
    CrearConvalidacionesCicloMasivasRequest,
    CrearConvalidacionCicloRequest,
    CrearModulosRequest,
)

ADMIN_UPLOAD_DIR = Path(os.getenv("FORM_UPLOAD_DIR", "/app/data/uploads")).resolve()
PUBLIC_ADMIN_PATHS = {"/admin/auth/config"}


def _normalized_request_path(request: Request) -> str:
    path = request.url.path.rstrip("/") or "/"
    root_path = (request.scope.get("root_path") or "").rstrip("/")
    if root_path and path.startswith(root_path):
        stripped = path[len(root_path):]
        return stripped or "/"
    return path


class AdminAuthRoute(APIRoute):
    def get_route_handler(self):
        original_route_handler = super().get_route_handler()

        async def custom_route_handler(request: Request):
            if _normalized_request_path(request) not in PUBLIC_ADMIN_PATHS:
                conn = database.connect()
                try:
                    admin_user = require_admin_user(request, conn)
                    conn.commit()
                finally:
                    conn.close()

                request.state.admin = admin_user

            return await original_route_handler(request)

        return custom_route_handler


router = APIRouter(
    prefix="/admin",
    tags=["admin"],
    route_class=AdminAuthRoute,
)


@router.get("/auth/config")
async def admin_auth_config():
    return {
        "enabled": True,
        "auth_path": "/auth",
    }


@router.get("/session")
async def admin_session(request: Request):
    admin_user = getattr(request.state, "admin", None)
    if not admin_user:
        raise HTTPException(status_code=401, detail="No hay sesión de administrador")
    return {"ok": True, **admin_user}


@router.post("/login")
async def login_admin():
    raise HTTPException(status_code=410, detail="El acceso admin ahora se gestiona desde /auth/login")


@router.post("/login/microsoft")
async def login_admin_microsoft():
    raise HTTPException(status_code=410, detail="El acceso admin ahora se gestiona desde /auth/login")


@router.get("/formularios")
async def listar_formularios(
    estado_id: Optional[int] = Query(default=None),
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Lista todos los formularios con datos básicos del alumno y solicitudes."""
    try:
        return database.AdminQueries.list_admin_formularios(db, estado_id=estado_id)
    except sqlite3.Error as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/documentos/abrir")
async def abrir_documento(path: str):
    """Sirve un documento aportado para que el admin pueda abrirlo."""
    try:
        file_path = Path(path).resolve()
    except Exception:
        raise HTTPException(status_code=400, detail="Ruta de documento inválida")

    if not str(file_path).startswith(str(ADMIN_UPLOAD_DIR)):
        raise HTTPException(status_code=403, detail="Ruta fuera del directorio permitido")
    if not file_path.is_file():
        raise HTTPException(status_code=404, detail="Documento no encontrado")

    return FileResponse(path=file_path)


@router.get("/exportar_solicitudes_convalidacion")
async def exportar_solicitudes_convalidacion(
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Exporta solicitudes de convalidación (formularios validados) en XLSX."""
    try:
        rows = database.AdminQueries.list_export_convalidaciones_rows(db)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Solicitudes"
        sheet.append(
            [
                "NOMBRE",
                "DNI",
                "EMAIL",
                "CICLO MATRICULADO",
                "MODULO A CONVALIDAR",
                "RESOL.",
                "CICLO CURSADO",
                "MOD CURSADO",
                "NOTA MOD.",
                "OBSERVACIONES",
            ]
        )
        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 16
        sheet.column_dimensions["C"].width = 34
        sheet.column_dimensions["D"].width = 30
        sheet.column_dimensions["E"].width = 34
        sheet.column_dimensions["F"].width = 12
        sheet.column_dimensions["G"].width = 34
        sheet.column_dimensions["H"].width = 42
        sheet.column_dimensions["I"].width = 12
        sheet.column_dimensions["J"].width = 36
        for row in rows:
            nota = row.get("nota_modulo")
            sheet.append(
                [
                    row.get("alumno_nombre") or "",
                    row.get("alumno_dni") or "",
                    row.get("alumno_email") or "",
                    row.get("ciclo_matriculado") or "",
                    row.get("modulo_a_convalidar") or "",
                    row.get("resolucion") or "",
                    row.get("ciclo_cursado") or "",
                    row.get("modulo_cursado") or "",
                    None if nota is None else int(nota),
                    row.get("observaciones") or "",
                ]
            )
            current_row = sheet.max_row
            sheet.cell(row=current_row, column=7).alignment = Alignment(wrap_text=True)
            sheet.cell(row=current_row, column=8).alignment = Alignment(wrap_text=True)

        output = io.BytesIO()
        workbook.save(output)
        content = output.getvalue()
        return Response(
            content=content,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": 'attachment; filename="solicitudes_convalidacion_validadas.xlsx"'
            },
        )
    except HTTPException:
        raise
    except sqlite3.Error as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.put("/formularios/{id_formulario}/estado")
async def cambiar_estado_formulario(
    id_formulario: int,
    request: CambiarEstadoFormularioRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Cambia el estado de un formulario a revision/validado/rechazado."""
    try:
        updated = database.AdminQueries.update_formulario_estado(
            db,
            id_formulario=id_formulario,
            estado_id=request.estado_id,
            admin_id=request.admin_id,
        )
        if updated == 0:
            raise HTTPException(status_code=404, detail="Formulario no encontrado")

        db.commit()
        return {"id": id_formulario, "estado_id": request.estado_id}
    except HTTPException:
        raise
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/formularios/{id_formulario}")
async def eliminar_formulario(
    id_formulario: int,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Elimina un formulario y sus datos relacionados (con borrado en cascada)."""
    try:
        archivos = database.AdminQueries.list_formulario_archivos(db, formulario_id=id_formulario)
        deleted = database.AdminQueries.delete_formulario(db, formulario_id=id_formulario)
        if deleted == 0:
            raise HTTPException(status_code=404, detail="Formulario no encontrado")

        db.commit()

        for archivo in archivos:
            ruta = (archivo.get("ruta_almacenamiento") or "").strip()
            if not ruta:
                continue

            try:
                file_path = Path(ruta).resolve()
            except Exception:
                continue

            if not str(file_path).startswith(str(ADMIN_UPLOAD_DIR)):
                continue

            file_path.unlink(missing_ok=True)

        return {"ok": True, "id": id_formulario}
    except HTTPException:
        raise
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.put("/solicitudes/{id_solicitud}/estado")
async def cambiar_estado_solicitud(
    id_solicitud: int,
    request: CambiarEstadoSolicitudRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Cambia el estado de una solicitud concreta a revision/validado/rechazado."""
    try:
        updated = database.AdminQueries.update_solicitud_estado(
            db,
            id_solicitud=id_solicitud,
            estado_modulo_id=request.estado_modulo_id,
            admin_id=request.admin_id,
            nota_manual=request.nota_manual,
        )
        if not updated:
            raise HTTPException(status_code=404, detail="Solicitud no encontrada")

        db.commit()
        return updated
    except HTTPException:
        raise
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/ciclos-modulos")
async def listar_ciclos_modulos(
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Lista todos los ciclos y sus módulos."""
    try:
        return database.AdminQueries.list_admin_ciclos_modulos(db)
    except sqlite3.Error as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/ciclos")
async def crear_ciclo(
    request: CrearCicloRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Crea un ciclo nuevo en catálogo."""
    try:
        nombre = (request.nombre or "").strip()
        if not nombre:
            raise HTTPException(status_code=400, detail="El nombre del ciclo es obligatorio")
        if request.id_familia <= 0 or request.id_grado <= 0:
            raise HTTPException(status_code=400, detail="Familia y grado son obligatorios")

        database.CatalogQueries.create_ciclo(
            db,
            nombre=nombre,
            id_familia=request.id_familia,
            id_grado=request.id_grado,
            es_somorrostro=1 if int(request.es_somorrostro) == 1 else 0,
        )
        db.commit()
        return {"ok": True}
    except HTTPException:
        raise
    except sqlite3.IntegrityError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.put("/ciclos/{id_ciclo}")
async def actualizar_ciclo(
    id_ciclo: int,
    request: ActualizarCicloRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Actualiza un ciclo existente en catálogo."""
    try:
        nombre = (request.nombre or "").strip()
        if not nombre:
            raise HTTPException(status_code=400, detail="El nombre del ciclo es obligatorio")
        if request.id_familia <= 0 or request.id_grado <= 0:
            raise HTTPException(status_code=400, detail="Familia y grado son obligatorios")

        updated = database.CatalogQueries.update_ciclo(
            db,
            ciclo_id=id_ciclo,
            nombre=nombre,
            id_familia=request.id_familia,
            id_grado=request.id_grado,
            es_somorrostro=1 if int(request.es_somorrostro) == 1 else 0,
        )
        if updated == 0:
            raise HTTPException(status_code=404, detail="Ciclo no encontrado")

        db.commit()
        return {"ok": True}
    except HTTPException:
        raise
    except sqlite3.IntegrityError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/convalidaciones")
async def listar_convalidaciones(
    grado_id: int | None = None,
    ciclo_id: int | None = None,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Lista todas las reglas de convalidación con módulo destino y origen."""
    try:
        return database.AdminQueries.list_admin_convalidaciones(
            db,
            grado_id=grado_id,
            ciclo_id=ciclo_id,
        )
    except sqlite3.Error as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/usuarios")
async def listar_usuarios(
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Lista usuarios registrados y su rol actual."""
    try:
        return database.UserQueries.list_users(db)
    except sqlite3.Error as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.put("/usuarios/{id_usuario}/rol")
async def actualizar_rol_usuario(
    id_usuario: int,
    request: ActualizarUsuarioRolRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Actualiza el rol de un usuario existente."""
    try:
        rol = (request.rol or "").strip().lower()
        if rol not in {"admin", "alumno"}:
            raise HTTPException(status_code=400, detail="El rol debe ser 'admin' o 'alumno'")

        updated = database.UserQueries.update_user_role(db, user_id=id_usuario, rol=rol)
        if updated == 0:
            raise HTTPException(status_code=404, detail="Usuario no encontrado")

        db.commit()
        return {
            "ok": True,
            "id": id_usuario,
            "rol": rol,
        }
    except HTTPException:
        raise
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/crear_convalidaciones")
async def crear_convalidacion(
    request: CrearConvalidacionRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Crea una regla de convalidación y sus módulos de origen."""
    try:
        id_modulo_destino = int(request.id_modulo_destino)
        if id_modulo_destino <= 0:
            raise HTTPException(status_code=400, detail="El módulo destino es obligatorio")

        source_link = (request.source_link or "").strip() or None
        source_page = int(request.source_page) if request.source_page is not None else None

        created = database.AdminQueries.create_convalidacion_rule(
            db,
            id_modulo_destino=id_modulo_destino,
            id_modulos_origen=request.id_modulos_origen or [],
            source_link=source_link,
            source_page=source_page,
        )

        db.commit()
        return {
            "ok": True,
            "id": int(created["id"]),
            "id_modulo_destino": int(created["id_modulo_destino"]),
            "id_modulos_origen": list(created["id_modulos_origen"]),
        }
    except HTTPException:
        raise
    except ValueError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.IntegrityError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/crear_convalidaciones_masivas")
async def crear_convalidaciones_masivas(
    request: CrearConvalidacionesMasivasRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Crea varias reglas de convalidación módulo a módulo en una sola petición."""
    try:
        source_link = (request.source_link or "").strip() or None
        source_page = int(request.source_page) if request.source_page is not None else None

        created = database.AdminQueries.create_convalidacion_rules_batch(
            db,
            reglas=[
                {
                    "id_modulo_destino": int(item.id_modulo_destino),
                    "id_modulo_origen": int(item.id_modulo_origen),
                }
                for item in (request.reglas or [])
            ],
            source_link=source_link,
            source_page=source_page,
        )

        db.commit()
        return {
            "ok": True,
            "created_count": int(created["created_count"]),
            "skipped_count": int(created["skipped_count"]),
            "skipped_existing_count": int(created.get("skipped_existing_count", 0)),
            "created": [
                {
                    "id": int(item["id"]),
                    "id_modulo_destino": int(item["id_modulo_destino"]),
                    "id_modulos_origen": list(item["id_modulos_origen"]),
                }
                for item in created["created"]
            ],
        }
    except HTTPException:
        raise
    except ValueError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.IntegrityError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/convalidaciones_masivas")
async def eliminar_convalidaciones_masivas(
    request: EliminarConvalidacionesMasivasRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Elimina varias reglas de convalidación módulo a módulo en una sola petición."""
    try:
        deleted = database.AdminQueries.delete_convalidacion_rules_batch(
            db,
            reglas=[
                {
                    "id_modulo_destino": int(item.id_modulo_destino),
                    "id_modulo_origen": int(item.id_modulo_origen),
                }
                for item in (request.reglas or [])
            ],
        )

        db.commit()
        return {
            "ok": True,
            "deleted_count": int(deleted["deleted_count"]),
            "skipped_count": int(deleted["skipped_count"]),
            "skipped_missing_count": int(deleted.get("skipped_missing_count", 0)),
        }
    except HTTPException:
        raise
    except ValueError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.IntegrityError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/crear_convalidaciones_ciclo")
async def crear_convalidacion_ciclo(
    request: CrearConvalidacionCicloRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Crea una regla de convalidación por ciclo completo."""
    try:
        id_modulo_destino = int(request.id_modulo_destino)
        id_ciclo_origen = int(request.id_ciclo_origen)
        if id_modulo_destino <= 0:
            raise HTTPException(status_code=400, detail="El módulo destino es obligatorio")
        if id_ciclo_origen <= 0:
            raise HTTPException(status_code=400, detail="El ciclo origen es obligatorio")

        source_link = (request.source_link or "").strip() or None
        source_page = int(request.source_page) if request.source_page is not None else None

        created = database.AdminQueries.create_convalidacion_ciclo_rule(
            db,
            id_modulo_destino=id_modulo_destino,
            id_ciclo_origen=id_ciclo_origen,
            source_link=source_link,
            source_page=source_page,
        )

        db.commit()
        return {
            "ok": True,
            "id": int(created["id"]),
            "id_modulo_destino": int(created["id_modulo_destino"]),
            "id_ciclo_origen": int(created["id_ciclo_origen"]),
        }
    except HTTPException:
        raise
    except ValueError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.IntegrityError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/crear_convalidaciones_ciclo_masivas")
async def crear_convalidaciones_ciclo_masivas(
    request: CrearConvalidacionesCicloMasivasRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    try:
        source_link = (request.source_link or "").strip() or None
        source_page = int(request.source_page) if request.source_page is not None else None

        created = database.AdminQueries.create_convalidacion_ciclo_rules_batch(
            db,
            reglas=[
                {
                    "id_modulo_destino": int(item.id_modulo_destino),
                    "id_ciclo_origen": int(item.id_ciclo_origen),
                }
                for item in (request.reglas or [])
            ],
            source_link=source_link,
            source_page=source_page,
        )

        db.commit()
        return {
            "ok": True,
            "created_count": int(created["created_count"]),
            "skipped_count": int(created["skipped_count"]),
            "skipped_existing_count": int(created.get("skipped_existing_count", 0)),
            "created": [
                {
                    "id": int(item["id"]),
                    "id_modulo_destino": int(item["id_modulo_destino"]),
                    "id_modulos_origen": [],
                }
                for item in created["created"]
            ],
        }
    except ValueError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/convalidaciones_ciclo_masivas")
async def eliminar_convalidaciones_ciclo_masivas(
    request: EliminarConvalidacionesCicloMasivasRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    try:
        deleted = database.AdminQueries.delete_convalidacion_ciclo_rules_batch(
            db,
            reglas=[
                {
                    "id_modulo_destino": int(item.id_modulo_destino),
                    "id_ciclo_origen": int(item.id_ciclo_origen),
                }
                for item in (request.reglas or [])
            ],
        )

        db.commit()
        return {
            "ok": True,
            "deleted_count": int(deleted["deleted_count"]),
            "skipped_count": int(deleted["skipped_count"]),
            "skipped_missing_count": int(deleted.get("skipped_missing_count", 0)),
        }
    except ValueError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/ciclos/{id_ciclo}")
async def eliminar_ciclo(
    id_ciclo: int,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Elimina un ciclo y sus datos relacionados (con borrado en cascada)."""
    try:
        deleted = database.CatalogQueries.delete_ciclo(db, id_ciclo)
        if deleted == 0:
            raise HTTPException(status_code=404, detail="Ciclo no encontrado")
        db.commit()
        return {"ok": True, "id": id_ciclo}
    except HTTPException:
        raise
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/modulos/{id_modulo}")
async def eliminar_modulo(
    id_modulo: int,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Elimina un módulo y sus datos relacionados (con borrado en cascada)."""
    try:
        deleted = database.CatalogQueries.delete_modulo(db, id_modulo)
        if deleted == 0:
            raise HTTPException(status_code=404, detail="Módulo no encontrado")
        db.commit()
        return {"ok": True, "id": id_modulo}
    except HTTPException:
        raise
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/convalidaciones/{id_convalidacion}")
async def eliminar_convalidacion(
    id_convalidacion: int,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Elimina una regla de convalidación por ID (y sus orígenes en cascada)."""
    try:
        deleted = database.AdminQueries.delete_convalidacion_rule(db, id_convalidacion)
        if deleted == 0:
            raise HTTPException(status_code=404, detail="Regla de convalidación no encontrada")
        db.commit()
        return {"ok": True, "id": id_convalidacion}
    except HTTPException:
        raise
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/convalidaciones-ciclo/{id_convalidacion_ciclo}")
async def eliminar_convalidacion_ciclo(
    id_convalidacion_ciclo: int,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Elimina una regla de convalidación de ciclo completo por ID."""
    try:
        deleted = database.AdminQueries.delete_convalidacion_ciclo_rule(db, id_convalidacion_ciclo)
        if deleted == 0:
            raise HTTPException(status_code=404, detail="Regla de convalidación por ciclo no encontrada")
        db.commit()
        return {"ok": True, "id": id_convalidacion_ciclo}
    except HTTPException:
        raise
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/convalidaciones/{id_convalidacion}/origenes/{id_modulo_origen}")
async def eliminar_convalidacion_origen(
    id_convalidacion: int,
    id_modulo_origen: int,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Elimina un módulo origen concreto de una regla de convalidación."""
    try:
        deleted = database.AdminQueries.delete_convalidacion_origen(
            db,
            convalidacion_id=id_convalidacion,
            modulo_origen_id=id_modulo_origen,
        )
        if deleted == 0:
            raise HTTPException(status_code=404, detail="Origen de convalidación no encontrado")
        db.commit()
        return {"ok": True, "id_convalidacion": id_convalidacion, "id_modulo_origen": id_modulo_origen}
    except HTTPException:
        raise
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/ciclos/{id_ciclo}/modulos")
async def crear_modulos(
    id_ciclo: int,
    request: CrearModulosRequest,
    db: sqlite3.Connection = Depends(database.get_db),
):
    """Crea uno o varios módulos en un ciclo existente."""
    try:
        if not database.CatalogQueries.exists_ciclo(db, id_ciclo):
            raise HTTPException(status_code=404, detail="Ciclo no encontrado")

        payload = [
            {
                "id_oficial": item.id_oficial,
                "nombre": item.nombre,
                "numerico": item.numerico,
            }
            for item in request.modulos
            if (item.nombre or "").strip()
        ]
        if not payload:
            raise HTTPException(status_code=400, detail="Debes enviar al menos un módulo con nombre")

        inserted = database.CatalogQueries.create_modulos(db, id_ciclo=id_ciclo, modulos=payload)
        if inserted == 0:
            raise HTTPException(status_code=400, detail="No se pudo crear ningún módulo")

        db.commit()
        return {"ok": True, "inserted": inserted}
    except HTTPException:
        raise
    except sqlite3.IntegrityError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except sqlite3.Error as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))
